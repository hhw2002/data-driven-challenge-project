"""Filter REE transmission-node access capacity (raw) into per-region clean CSVs.

REE publishes one workbook for demand and one for generation, each with one
row per transmission-network node (220/400 kV) and a multi-row header built
from merged cells. This flattens that header into single column names
(levels joined with " | "), keeps every column, and writes the region's
rows to data/grid/clean/.

Only the "Ficha Capacidades" sheet is used. The "Zonas" sheet (nodes that
share capacity as a zone) and "Leyenda" (column definitions) are left in
the raw workbook.
"""
from pathlib import Path

import openpyxl
import pandas as pd

PROJECT_DIR = Path(__file__).resolve().parent.parent
RAW_DIR = PROJECT_DIR / "data" / "grid" / "raw" / "ree_capacidad_transporte"
CLEAN_DIR = PROJECT_DIR / "data" / "grid" / "clean"

# file names start with the month the data describes, not the publication month
# (REE's own name for these is 2026_10_01_GRT_*.xlsx)
FILES = {
    "demanda": RAW_DIR / "2026_09_demanda_ree.xlsx",
    "generacion": RAW_DIR / "2026_09_generacion_ree.xlsx",
}

SHEET = "Ficha Capacidades"

REGIONS = {
    "catalunya": "Cataluña",
}

NODE_COL = "Nombre y tensión del nudo"
# placeholders REE uses for "not applicable" in otherwise numeric columns
NOT_APPLICABLE = {"NA", "N/A"}


def _clean(value) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def read_sheet(path: Path) -> pd.DataFrame:
    """Read the capacity sheet with its merged multi-row header flattened."""
    ws = openpyxl.load_workbook(path, data_only=True)[SHEET]
    grid = [[_clean(v) for v in row] for row in ws.iter_rows(values_only=True)]

    # the node-name header cell is merged down the full height of the header
    n_header = next(
        r.max_row for r in ws.merged_cells.ranges if r.min_row == 1 and r.min_col == 1
    )

    # a merged cell only stores its value top-left: copy it across the range
    for rng in ws.merged_cells.ranges:
        if rng.min_row > n_header:
            continue
        value = grid[rng.min_row - 1][rng.min_col - 1]
        for r in range(rng.min_row - 1, min(rng.max_row, n_header)):
            for c in range(rng.min_col - 1, rng.max_col):
                grid[r][c] = value

    columns = []
    for c in range(ws.max_column):
        levels = []
        for r in range(n_header):
            if grid[r][c] and grid[r][c] not in levels:
                levels.append(grid[r][c])
        columns.append(" | ".join(levels))

    df = pd.DataFrame(grid[n_header:], columns=columns)
    df = df[df[NODE_COL] != ""].reset_index(drop=True)
    return df.apply(_to_numeric_if_possible)


def _to_numeric_if_possible(col: pd.Series) -> pd.Series:
    """Make a column numeric when all it holds is numbers and N/A markers."""
    filled = col[(col != "") & ~col.isin(NOT_APPLICABLE)]
    if filled.empty:
        return col
    numeric = pd.to_numeric(filled, errors="coerce")
    if numeric.isna().any():
        return col
    return pd.to_numeric(col, errors="coerce")


def add_node_parts(df: pd.DataFrame) -> pd.DataFrame:
    """Split 'PIEROLA 400' into nombre_nudo='PIEROLA' and tension_kv=400."""
    parts = df[NODE_COL].str.rsplit(" ", n=1, expand=True)
    df.insert(1, "nombre_nudo", parts[0])
    df.insert(2, "tension_kv", pd.to_numeric(parts[1]))
    return df


def main() -> None:
    for kind, path in FILES.items():
        df = add_node_parts(read_sheet(path))
        for region_key, region_label in REGIONS.items():
            CLEAN_DIR.mkdir(parents=True, exist_ok=True)
            subset = df[df["Comunidad Autónoma"] == region_label]
            out_path = CLEAN_DIR / f"ree_{kind}_{region_key}.csv"
            subset.to_csv(out_path, index=False)
            print(f"{region_key}/ree_{kind}: {len(subset)} nodes, {len(subset.columns)} columns -> {out_path}")


if __name__ == "__main__":
    main()
